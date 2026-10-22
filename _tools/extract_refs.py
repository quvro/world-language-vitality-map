# -*- coding: utf-8 -*-
"""Dump text from the teacher's reference PPTX and DOCX into plain-text files."""
import os
import sys

from pptx import Presentation
from docx import Document

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "_refs_out")
os.makedirs(OUT, exist_ok=True)


def dump_pptx(path: str, out_path: str) -> None:
    prs = Presentation(path)
    lines: list[str] = []
    lines.append(f"# FILE: {path}")
    lines.append(f"# slides: {len(prs.slides)}")
    for idx, slide in enumerate(prs.slides, 1):
        lines.append("")
        lines.append(f"===== SLIDE {idx} =====")
        for shape in slide.shapes:
            stype = shape.shape_type
            name = shape.name
            if shape.has_text_frame:
                text = "\n".join(
                    p.text for p in shape.text_frame.paragraphs
                ).strip()
                if text:
                    lines.append(f"[{name} | {stype}] {text}")
            elif shape.has_table:
                lines.append(f"[TABLE | {name}]")
                for row in shape.table.rows:
                    cells = [c.text.replace("\n", " ").strip() for c in row.cells]
                    lines.append("  | " + " | ".join(cells) + " |")
            else:
                lines.append(f"[{stype} | {name}]")
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append(f"[NOTES] {notes}")
    with open(out_path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))
    print(f"pptx -> {out_path}  ({len(prs.slides)} slides, {len(lines)} lines)")


def dump_docx(path: str, out_path: str) -> None:
    doc = Document(path)
    lines: list[str] = []
    lines.append(f"# FILE: {path}")

    body = doc.element.body
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in body.iterchildren():
        tag = child.tag.split("}")[-1]
        if tag == "p":
            para = Paragraph(child, doc)
            text = para.text.strip()
            if text:
                style = para.style.name if para.style is not None else ""
                prefix = f"<{style}> " if style and style != "Normal" else ""
                lines.append(prefix + text)
        elif tag == "tbl":
            table = Table(child, doc)
            lines.append("[TABLE]")
            for row in table.rows:
                cells = [c.text.replace("\n", " ").strip() for c in row.cells]
                lines.append("  | " + " | ".join(cells) + " |")
            lines.append("[/TABLE]")

    # headers / footers
    for i, section in enumerate(doc.sections):
        for label, part in (("header", section.header), ("footer", section.footer)):
            txt = "\n".join(p.text for p in part.paragraphs).strip()
            if txt:
                lines.append(f"[SECTION {i} {label.upper()}] {txt}")

    with open(out_path, "w", encoding="utf-8") as fh:
        fh.write("\n".join(lines))
    print(f"docx -> {out_path}  ({len(doc.paragraphs)} paras, {len(lines)} lines)")


if __name__ == "__main__":
    # 原来这里硬编码了作者本机的两个路径（在桌面上），
    # 既不可移植，也会把私人路径带进版本库。改成命令行传入。
    # 用法：
    #   python extract_refs.py <输入.pptx|输入.docx> [输出目录]
    if len(sys.argv) < 2:
        print("用法: python extract_refs.py <输入.pptx|输入.docx> [输出目录]")
        print("作用: 把老师给的参考 PPTX / DOCX 抽成纯文本，便于检索需求对齐点。")
        raise SystemExit(2)
    src = sys.argv[1]
    out_dir = sys.argv[2] if len(sys.argv) > 2 else OUT
    os.makedirs(out_dir, exist_ok=True)
    stem = os.path.splitext(os.path.basename(src))[0]
    dest = os.path.join(out_dir, stem + ".txt")
    if src.lower().endswith(".pptx"):
        dump_pptx(src, dest)
    elif src.lower().endswith(".docx"):
        dump_docx(src, dest)
    else:
        print(f"不支持的文件类型: {src}")
        raise SystemExit(2)
